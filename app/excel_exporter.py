from __future__ import annotations
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple, TypedDict, cast
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload
import textwrap

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import RowDimension
from openpyxl.worksheet.worksheet import Worksheet
from .scheduler import build_timetable_payload, teacher_load_rows, blend_hex_colors
from .models import Group, GroupStudyProgram, StudyProgram, ScheduledClass

def default_font(
    *,
    name: str = 'Times New Roman',
    size: Optional[int] = None,
    bold: bool = False,
    italic: bool = False,
    underline: Optional[str] = None,
    color: Optional[str] = None,
) -> Font:
    kwargs: Dict[str, Any] = {'name': name, 'bold': bold, 'italic': italic}
    if size is not None:
        kwargs['size'] = size
    if underline is not None:
        kwargs['underline'] = underline
    if color is not None:
        kwargs['color'] = color
    return Font(**kwargs)


import re

FILL_MAP = {
    "ielts": "9FC5E8",
    "german": "F9CB9C",
    "ae": "EAD1DC",
    "core": "F4B183",
    "shared": "D99A9A",
    "elective": "C49A00",
    "other": "D9D2E9",
}

PROGRAM_FILL_VARIANTS = [
    "FFF2CC",
    "E8F0D9",
    "D9E8F8",
    "F9E2E6",
    "EDE7F5",
    "F7EED9",
    "E8F2E8",
    "F8E7F2",
    "DFF0EB",
    "FAE9D3",
    "E9E8F3",
    "F3F0E8",
    "DDE8F0",
    "F8ECEA",
    "E9F1EF",
    "F8F2DA",
    "EDE9EC",
    "DDE9E4",
    "FDF3D8",
    "E8E8F0",
]

TAG_FILL_VARIANTS = [
    "F4B183",
    "B6D7A8",
    "C9DAF8",
    "EAD1DC",
    "FCE4D6",
    "D0E0E3",
    "FFF2CC",
    "D9D2E9",
    "E2EFDA",
    "F9CB9C",
]

BLANK_FILL = "D9D9D9"
REQUIRED_ROW_HEIGHT = 32
OVERLAY_ROW_HEIGHT = 40
_LINE_HEIGHT_PTS = 15
_OVERLAY_MIN_HEIGHT = 18


def _format_item(item: Dict[str, Any]) -> str:
    course_name = str(item["course_name"])
    parts: List[str] = []
    if item.get("kind") == "elective":
        parts.append("(Elective)")
    parts.append(course_name)
    group_codes: List[str] = cast(List[str], item.get("group_codes") or [])
    if item.get("kind") == "required":
        if len(group_codes) > 1:
            parts.append(f"({', '.join(group_codes)})")
    elif item.get("kind") != "elective" and group_codes:
        parts.append(f"({', '.join(group_codes)})")
    if item.get("all_group"):
        parts.append("(all group)")
    if item["program_codes"]:
        parts.append(f"[{', '.join(item['program_codes'])}]")
    if item["notes"]:
        parts.append(item["notes"])
    if item["teacher_name"]:
        parts.append(item["teacher_name"])
    if item["room_name"]:
        parts.append(f"Room: {item['room_name']}")
    if item.get("kind") == "required":
        return "\n".join(parts)
    return " ".join(parts)


def _format_german_item(item: Dict[str, Any]) -> str:
    """Format German class blocks without repeating the course name."""
    parts: List[str] = []
    if item.get("kind") == "elective":
        parts.append("(Elective)")
    group_codes: List[str] = cast(List[str], item.get("group_codes") or [])
    if item.get("kind") == "required":
        if len(group_codes) > 1:
            parts.append(f"({', '.join(group_codes)})")
    elif item.get("kind") != "elective" and group_codes:
        parts.append(f"({', '.join(group_codes)})")
    if item.get("all_group"):
        parts.append("(all group)")
    if item["program_codes"]:
        parts.append(f"[{', '.join(item['program_codes'])}]")
    if item["notes"]:
        parts.append(item["notes"])
    if item["teacher_name"]:
        parts.append(item["teacher_name"])
    if item["room_name"]:
        parts.append(f"Room: {item['room_name']}")
    if item.get("kind") == "required":
        return "\n".join(parts)
    return " ".join(parts)


def _german_group_summary(timeslots: List[Dict[str, Any]], cell_map: Dict[str, Any], groups: List[Dict[str, Any]]) -> Dict[int, str]:
    counts: Dict[int, Counter[str]] = {group['id']: Counter() for group in groups}
    for slot in timeslots:
        if not str(slot.get("label", "")).startswith("German"):
            continue
        slot_items = cell_map.get(str(slot["id"]), {})
        for group in groups:
            items = slot_items.get(str(group["id"]), [])
            for item in items:
                course_name = str(item.get("course_name") or "").strip()
                if course_name:
                    counts[group['id']][course_name] += 1
    result: Dict[int, str] = {}
    for group_id, counter in counts.items():
        if not counter:
            result[group_id] = ""
            continue
        common, _ = sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))[0]
        result[group_id] = common
    return result


def _estimate_line_count(text: str, width_cols: int) -> int:
    max_chars = int(width_cols) if width_cols > 10 else max(12, int(24 * width_cols))
    total_lines = 0
    paragraphs = text.split("\n")

    for paragraph in paragraphs:
        if paragraph.strip() == "":
            total_lines += 1
            continue

        wrapped = textwrap.wrap(paragraph, width=max_chars, break_long_words=True, break_on_hyphens=False)
        total_lines += max(1, len(wrapped))

    return max(1, total_lines)




def _column_range_width_chars(ws: Worksheet, start_col: int, end_col: int) -> int:
    total_width = 0.0
    for col_idx in range(start_col, end_col + 1):
        letter = get_column_letter(col_idx)
        width = getattr(ws.column_dimensions[letter], "width", None)
        total_width += float(width) if width not in (None, 0) else 10.0
    # Use a slightly smaller effective width to ensure enough wrap height.
    return int(max(8, total_width * 0.88))


def _row_height_for_text(text: str, width_cols: int = 1, min_height: int = 24, font_size: int = 11) -> int:
    effective_width = width_cols
    if effective_width < 10:
        effective_width = max(8, int(effective_width * 24))
    estimated_chars = max(8, int(effective_width * 0.88))
    line_count = _estimate_line_count(text, estimated_chars)
    line_height = max(font_size * 1.35, float(_LINE_HEIGHT_PTS))
    return max(min_height, int(line_count * line_height + 10))


def _normalize_program_codes(program_codes: List[str]) -> List[str]:
    return sorted({str(code).strip() for code in program_codes if str(code).strip()})


def _program_base_color(program_code: str) -> str:
    program_code = str(program_code).strip()
    if not program_code:
        return PROGRAM_FILL_VARIANTS[0]
    index = abs(hash(program_code)) % len(PROGRAM_FILL_VARIANTS)
    return PROGRAM_FILL_VARIANTS[index]


def _program_fill_color(program_codes: List[str]) -> str:
    program_codes = _normalize_program_codes(program_codes)
    if not program_codes:
        return PROGRAM_FILL_VARIANTS[0]
    program_colors = [f"#{_program_base_color(code)}" for code in program_codes]
    if len(program_colors) == 1:
        return program_colors[0][1:]
    return blend_hex_colors(program_colors)


def _normalize_excel_color(value: Optional[str]) -> str:
    if not value:
        return "FFFFFFFF"
    color = str(value).strip()
    if color.startswith("#"):
        color = color[1:]
    if len(color) == 3:
        color = ''.join(ch * 2 for ch in color)
    color = color.upper()
    if re.fullmatch(r'[0-9A-F]{6}', color):
        return f"FF{color}"
    if re.fullmatch(r'[0-9A-F]{8}', color):
        return color
    return "FFFFFFFF"


def _get_fill_color(item: Dict[str, Any]) -> str:
    if item.get("fill_color"):
        return _normalize_excel_color(item["fill_color"])
    if item.get("kind") == "required":
        #get color based on course tag if available, otherwise use default required color
        if item.get("course_tag_fill_color"):
            return _normalize_excel_color(item["course_tag_fill_color"])
        
    if item.get("kind") == "program":
        program_codes = _normalize_program_codes(item.get("program_codes", []))
        return _program_fill_color(program_codes)

    if item.get("kind") == "elective":
        return FILL_MAP.get("elective", "C49A00")

    color_key = item.get("color_key") or str(item.get("course_code", "other"))
    return FILL_MAP.get(color_key, TAG_FILL_VARIANTS[abs(hash(color_key)) % len(TAG_FILL_VARIANTS)])


def _assign_program_colors_for_slot(overlays: List[Dict[str, Any]], program_color_map: Dict[str, str]) -> None:
    used: set[str] = set(program_color_map.values())
    used.update(FILL_MAP.values())
    used.update(TAG_FILL_VARIANTS)
    for item in overlays:
        if item.get("kind") != "program":
            continue
        if item.get("fill_color"):
            continue
        program_codes = _normalize_program_codes(item.get("program_codes", []))
        if not program_codes:
            continue
        program_key = "|".join(program_codes)
        if program_key not in program_color_map:
            blended_color = _program_fill_color(program_codes)
            program_color_map[program_key] = blended_color
            used.add(blended_color)
        item["fill_color"] = program_color_map[program_key]


def _build_program_color_map(payload: Dict[str, Any]) -> Dict[str, str]:
    program_color_map: Dict[str, str] = {}
    groups = payload["groups"]
    timeslots = sorted(payload["timeslots"], key=lambda t: t["sort_order"])
    weekday_order = sorted({(t["day_index"], t["weekday"]) for t in timeslots}, key=lambda d: d[0])
    timeslots_by_day = [ts for day_index, _ in weekday_order for ts in sorted([t for t in timeslots if t["day_index"] == day_index], key=lambda t: t["sort_order"])]

    for slot in timeslots_by_day:
        slot_items = payload["cells"].get(str(slot["id"]), {})
        overlay_rows_content: List[List[Dict[str, Any]]] = []
        for group in groups:
            items = _overlay_items(slot_items.get(str(group["id"]), []))
            deduped: List[Dict[str, Any]] = []
            for item in items:
                if not any(_same_overlay(item, d) for d in deduped):
                    deduped.append(item)
            overlay_rows_content.append(deduped)

        unique_overlays: List[Dict[str, Any]] = []
        for overlays in overlay_rows_content:
            for item in overlays:
                found = next((uo for uo in unique_overlays if _same_overlay_connectable(item, uo)), None)
                if found:
                    found["program_codes"] = _normalize_program_codes(found.get("program_codes", []) + item.get("program_codes", []))
                    found["group_codes"] = sorted(set(found.get("group_codes", []) + item.get("group_codes", [])))
                    found["all_group"] = found.get("all_group", False) or item.get("all_group", False)
                elif not any(_same_overlay(item, uo) for uo in unique_overlays):
                    unique_overlays.append(
                        {
                            **item,
                            "program_codes": _normalize_program_codes(item.get("program_codes", [])),
                            "group_codes": list(item.get("group_codes", [])),
                            "all_group": item.get("all_group", False),
                        }
                    )

        kind_order = {"required": 0, "program": 1, "elective": 2}
        unique_overlays.sort(
            key=lambda item: (
                kind_order.get(str(item.get("kind") or ""), 3),
                item.get("course_name", ""),
                item.get("teacher_name", ""),
                item.get("room_name", ""),
            )
        )
        _assign_program_colors_for_slot(unique_overlays, program_color_map)

    return program_color_map


def _required_item(items: List[Dict[str, Any]]) -> Dict[str, Any] | None:
    for item in items:
        if item["kind"] == "required":
            return item
    return None


def _overlay_items(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [item for item in items if item["kind"] != "required"]


def _same_overlay(a: Dict[str, Any], b: Dict[str, Any]) -> bool:
    keys = [
        "course_name",
        "teacher_name",
        "room_name",
        "notes",
        "kind",
        "color_key",
        "shared",
    ]
    return all(a.get(k) == b.get(k) for k in keys) and a.get("program_codes", []) == b.get("program_codes", [])


def _same_overlay_connectable(a: Dict[str, Any], b: Dict[str, Any]) -> bool:
    # Connect study-program overlays if they share the same course, teacher, room, and kind.
    return (
        a.get("kind") == "program"
        and b.get("kind") == "program"
        and a.get("course_name") == b.get("course_name")
        and a.get("teacher_name") == b.get("teacher_name")
        and a.get("room_name") == b.get("room_name")
        and a.get("notes") == b.get("notes")
        and a.get("color_key") == b.get("color_key")
    )


class _RowEntry(TypedDict):
    text: str
    fill_color: str
    col_span: int


def _cell(ws: Worksheet, row: int, column: int, value: Any | None = None) -> Cell:
    return cast(Cell, ws.cell(row=row, column=column, value=value))


def _parse_time_minutes(value: Any | None) -> Optional[int]:
    if value is None:
        return None
    text = str(value).strip().replace('.', ':')
    parts = [p.strip() for p in text.split(':') if p.strip()]
    if not parts:
        return None
    try:
        hours = int(parts[0])
        minutes = int(parts[1]) if len(parts) > 1 else 0
        return hours * 60 + minutes
    except ValueError:
        return None


def _write_overlay_merge(
    ws: Worksheet,
    row_idx: int,
    start_col: int,
    end_col: int,
    item: Dict[str, Any],
    border: Border,
) -> None:
    ws.merge_cells(
        start_row=row_idx,
        start_column=start_col,
        end_row=row_idx,
        end_column=end_col,
    )
    cell = _cell(ws, row_idx, start_col, _format_item(item))
    cell.fill = PatternFill("solid", fgColor=_get_fill_color(item))
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.font = default_font(bold=True, size=15)
    for c in range(start_col, end_col + 1):
        ws.cell(row=row_idx, column=c).border = border


def _write_group_timetable_sheet(
    ws: Worksheet,
    payload: Dict[str, Any],
    group: Dict[str, Any],
    blank_fill: PatternFill,
    border: Border,
    program_color_map: Dict[str, str],
) -> None:
    timeslots = sorted(payload["timeslots"], key=lambda t: t["sort_order"])
    group_id = group["id"]
    program_codes = [p["code"] for p in group.get("programs", []) if p.get("code")]
    title = f"Group {group['code']}"
    if program_codes:
        title += f" ({', '.join(program_codes)})"
    ws.title = _sanitize_sheet_title(title)

    timeslots_by_day: Dict[str, List[Dict[str, Any]]] = {}
    for ts in timeslots:
        timeslots_by_day.setdefault(ts["weekday"], []).append(ts)

    active_days: List[Tuple[int, str]] = sorted({(t["day_index"], t["weekday"]) for t in timeslots}, key=lambda d: d[0])
    active_day_slots: List[Dict[str, Any]] = [ts for ts in timeslots if (ts["day_index"], ts["weekday"]) in active_days]

    time_labels: List[str] = []
    label_to_slot: Dict[str, Dict[str, Any]] = {}
    for timeslot in sorted(active_day_slots, key=lambda t: (t["day_index"], t["sort_order"])):
        if timeslot["label"] not in time_labels:
            time_labels.append(timeslot["label"])
            label_to_slot[timeslot["label"]] = timeslot

    days = active_days

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(days))
    title_cell = _cell(ws, 1, 1, title)
    title_cell.font = default_font(bold=True, size=15)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    header_cell = ws.cell(row=2, column=1, value="Time")
    header_cell.font = default_font(bold=True, size=11)
    header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_cell.border = border
    ws.column_dimensions["A"].width = float(18)

    for idx, (_day_index, weekday) in enumerate(days, start=2):
        cell = ws.cell(row=2, column=idx, value=weekday)
        cell.font = default_font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(idx)].width = float(26)

    elective_items: List[Dict[str, Any]] = []
    row = 3
    for idx, label in enumerate(time_labels):
        day_cells: Dict[int, List[Dict[str, Any]]] = {}
        for day_index, weekday in days:
            day_slots = [ts for ts in timeslots if ts["day_index"] == day_index and ts["label"] == label]
            if not day_slots:
                day_cells[day_index] = []
                continue
            slot_items: List[Dict[str, Any]] = []
            for timeslot in day_slots:
                candidate_items = cast(List[Dict[str, Any]], payload["cells"].get(str(timeslot["id"]), {}).get(str(group_id), []))
                if candidate_items:
                    slot_items = candidate_items
                    break
            if not slot_items:
                timeslot = day_slots[0]
                slot_items = cast(List[Dict[str, Any]], payload["cells"].get(str(timeslot["id"]), {}).get(str(group_id), []))
            expanded_items: List[Dict[str, Any]] = []
            for item in slot_items:
                if item.get("kind") == "elective":
                    elective_items.append({
                        "day": weekday,
                        "time": label,
                        "item": item,
                    })
                    continue
                expanded_items.append(item)
            day_cells[day_index] = expanded_items

        rows_for_label = max(max(len(items), 1) for items in day_cells.values())
        if rows_for_label > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row + rows_for_label - 1, end_column=1)
        time_cell = _cell(ws, row, 1, label)
        time_cell.font = default_font(bold=True)
        time_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        time_cell.border = border

        for row_offset in range(rows_for_label):
            current_row = row + row_offset
            if row_offset > 0:
                merged_time_cell = _cell(ws, current_row, 1, None)
                merged_time_cell.border = border
            row_height = _OVERLAY_MIN_HEIGHT
            row_entries: List[Optional[_RowEntry]] = []
            row_items: List[Optional[Dict[str, Any]]] = []
            for day_index, weekday in days:
                items = day_cells.get(day_index, [])
                if row_offset >= len(items):
                    row_entries.append(None)
                    row_items.append(None)
                    continue

                item = items[row_offset]
                if item.get("kind") == "program":
                    program_codes = _normalize_program_codes(item.get("program_codes", []))
                    program_key = "|".join(program_codes)
                    if not item.get("fill_color") and program_key in program_color_map:
                        item["fill_color"] = program_color_map[program_key]
                text = _format_item(item)
                fill_color = _normalize_excel_color(_get_fill_color(item))
                row_entries.append({"text": text, "fill_color": fill_color, "col_span": 1})
                row_items.append(item)

            def _merge_overlay_segment(items: List[Dict[str, Any]]) -> _RowEntry:
                merged = dict(items[0])
                merged["program_codes"] = _normalize_program_codes(
                    [code for item in items for code in item.get("program_codes", [])]
                )
                merged["group_codes"] = sorted(
                    {code for item in items for code in item.get("group_codes", [])}
                )
                merged["all_group"] = any(item.get("all_group", False) for item in items)
                fill_colors: List[str] = []
                for item in items:
                    fill_color = item.get("fill_color")
                    if fill_color:
                        normalized = _normalize_excel_color(fill_color)
                        if normalized and len(normalized) == 8:
                            fill_colors.append(f"#{normalized[2:]}")
                if len(set(fill_colors)) > 1:
                    merged["fill_color"] = blend_hex_colors(fill_colors)
                else:
                    merged["fill_color"] = None
                merged_text = _format_item(merged)
                merged_fill = _normalize_excel_color(_get_fill_color(merged))
                return {"text": merged_text, "fill_color": merged_fill, "col_span": len(items)}

            segment_start: Optional[int] = None
            segment_items: List[Dict[str, Any]] = []
            for idx, item in enumerate(row_items):
                if item is None or item.get("merge_id") is None:
                    if segment_start is not None and len(segment_items) > 1:
                        merged_entry = _merge_overlay_segment(segment_items)
                        row_entries[segment_start] = merged_entry
                        for skip_idx in range(segment_start + 1, idx):
                            row_entries[skip_idx] = None
                    segment_start = None
                    segment_items = []
                    continue
                if segment_start is None:
                    segment_start = idx
                    segment_items = [item]
                elif item.get("merge_id") == segment_items[0].get("merge_id"):
                    segment_items.append(item)
                else:
                    if len(segment_items) > 1:
                        merged_entry = _merge_overlay_segment(segment_items)
                        row_entries[segment_start] = merged_entry
                        for skip_idx in range(segment_start + 1, idx):
                            row_entries[skip_idx] = None
                    segment_start = idx
                    segment_items = [item]
            if segment_start is not None and len(segment_items) > 1:
                merged_entry = _merge_overlay_segment(segment_items)
                row_entries[segment_start] = merged_entry
                for skip_idx in range(segment_start + 1, len(row_entries)):
                    row_entries[skip_idx] = None

            col_idx = 2
            while col_idx < 2 + len(row_entries):
                entry = row_entries[col_idx - 2]
                cell = _cell(ws, current_row, col_idx)
                cell.border = border
                if entry is None:
                    cell.fill = blank_fill
                    col_idx += 1
                    continue

                assert entry is not None
                span = entry["col_span"]
                end_col = col_idx + span - 1
                if span > 1:
                    ws.merge_cells(start_row=current_row, start_column=col_idx, end_row=current_row, end_column=end_col)
                cell = _cell(ws, current_row, col_idx)
                cell.value = entry["text"]
                cell.fill = PatternFill("solid", fgColor=str(entry["fill_color"]))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = default_font(bold=True, size=12)
                for c in range(col_idx, end_col + 1):
                    ws.cell(row=current_row, column=c).border = border
                    if c != col_idx:
                        _cell(ws, current_row, c).value = None
                column_width = _column_range_width_chars(ws, col_idx, end_col)
                needed_height = _row_height_for_text(entry["text"], width_cols=column_width, min_height=_OVERLAY_MIN_HEIGHT, font_size=12)
                row_height = max(row_height, needed_height)
                col_idx = end_col + 1

            if label.startswith("German") and len(days) > 1:
                segment_start = None
                segment_entry = None
                for idx_day, entry in enumerate(row_entries):
                    if entry is None:
                        if segment_start is not None and idx_day - segment_start > 1:
                            assert segment_entry is not None
                            start_col = 2 + segment_start
                            end_col = 2 + idx_day - 1
                            ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
                            merged_cell = _cell(ws, current_row, start_col)
                            merged_cell.value = segment_entry["text"]
                            merged_cell.fill = PatternFill("solid", fgColor=str(segment_entry["fill_color"]))
                            merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            merged_cell.font = default_font(bold=True, size=12)
                            for c in range(start_col, end_col + 1):
                                ws.cell(row=current_row, column=c).border = border
                                if c != start_col:
                                    _cell(ws, current_row, c).value = None
                        segment_start = None
                        segment_entry = None
                        continue

                    if segment_start is None:
                        segment_start = idx_day
                        segment_entry = entry
                        continue

                    assert entry is not None
                    if segment_entry is not None and (entry["text"] != segment_entry["text"] or entry["fill_color"] != segment_entry["fill_color"]):
                        if idx_day - segment_start > 1:
                            segment_entry_value = segment_entry
                            start_col = 2 + segment_start
                            end_col = 2 + idx_day - 1
                            ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
                            merged_cell = _cell(ws, current_row, start_col)
                            merged_cell.value = segment_entry_value["text"]
                            merged_cell.fill = PatternFill("solid", fgColor=str(segment_entry_value["fill_color"]))
                            merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            merged_cell.font = default_font(bold=True, size=12)
                            for c in range(start_col, end_col + 1):
                                ws.cell(row=current_row, column=c).border = border
                                if c != start_col:
                                    _cell(ws, current_row, c).value = None
                        segment_start = idx_day
                        segment_entry = entry

                if segment_start is not None and len(row_entries) - segment_start > 1:
                    assert segment_entry is not None
                    segment_entry_value = segment_entry
                    start_col = 2 + segment_start
                    end_col = 2 + len(row_entries) - 1
                    ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
                    merged_cell = _cell(ws, current_row, start_col)
                    merged_cell.value = segment_entry_value["text"]
                    merged_cell.fill = PatternFill("solid", fgColor=str(segment_entry_value["fill_color"]))
                    merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    merged_cell.font = default_font(bold=True, size=12)
                    for c in range(start_col, end_col + 1):
                        ws.cell(row=current_row, column=c).border = border
                        if c != start_col:
                            _cell(ws, current_row, c).value = None

            row_dimension = ws.row_dimensions[current_row]
            current_height = float(getattr(row_dimension, "height", 0) or 0)
            row_dimension.height = float(max(current_height, row_height))

        row += rows_for_label
        if idx < len(time_labels) - 1:
            current_slot = label_to_slot.get(label)
            next_slot = label_to_slot.get(time_labels[idx + 1])
            current_end = _parse_time_minutes(current_slot.get("end") if current_slot else None)
            next_start = _parse_time_minutes(next_slot.get("start") if next_slot else None)
            if current_end is not None and next_start is not None and current_end <= 12 * 60 and next_start >= 13 * 60:
                lunch_row = row
                ws.merge_cells(start_row=lunch_row, start_column=1, end_row=lunch_row, end_column=1 + len(days))
                lunch_cell = _cell(ws, lunch_row, 1, "Lunch break")
                lunch_cell.fill = PatternFill("solid", fgColor="F2F2F2")
                lunch_cell.alignment = Alignment(horizontal="center", vertical="center")
                lunch_cell.font = default_font(italic=True)
                for c in range(1, 2 + len(days)):
                    ws.cell(row=lunch_row, column=c).border = border
                row += 1

    if elective_items:
        row += 1
        header_row = row
        ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row, end_column=1 + len(days))
        header_cell = _cell(ws, header_row, 1, "Elective classes")
        header_cell.font = default_font(bold=True, size=12)
        header_cell.alignment = Alignment(horizontal="left", vertical="center")
        row += 1

        ws.cell(row=row, column=1, value="Day").font = default_font(bold=True)
        ws.cell(row=row, column=2, value="Time").font = default_font(bold=True)
        ws.cell(row=row, column=3, value="Class details").font = default_font(bold=True)
        ws.row_dimensions[row].height = 24
        row += 1

        for elective in elective_items:
            ws.cell(row=row, column=1, value=elective["day"]) .border = border
            ws.cell(row=row, column=2, value=elective["time"]) .border = border
            detail = _format_item(elective["item"])
            detail_cell = _cell(ws, row, 3, detail)
            detail_cell.alignment = Alignment(wrap_text=True, vertical="center")
            detail_cell.border = border
            detail_cell.font = default_font(bold=True, size=12)
            detail_width = _column_range_width_chars(ws, 3, 3)
            ws.row_dimensions[row].height = _row_height_for_text(detail, width_cols=detail_width, min_height=_OVERLAY_MIN_HEIGHT, font_size=12)
            row += 1


def _build_teacher_schedule_rows(
    db: Session,
    timetable_id: Optional[int],
    teacher_id: int,
    group_ids: Optional[List[int]] = None,
    study_program_ids: Optional[List[int]] = None,
) -> List[Dict[str, Any]]:
    query = select(ScheduledClass).options(
        joinedload(ScheduledClass.group),
        joinedload(ScheduledClass.timeslot),
        joinedload(ScheduledClass.room),
        joinedload(ScheduledClass.course),
        joinedload(ScheduledClass.study_program),
    ).where(
        ScheduledClass.deploy.is_(True),
        ScheduledClass.teacher_id == teacher_id,
    )
    if timetable_id is not None:
        query = query.join(Group).where(Group.timetable_id == timetable_id)
    if group_ids is not None:
        query = query.where(ScheduledClass.group_id.in_(group_ids))
    classes = db.execute(query).unique().scalars().all()

    grouped: Dict[Tuple[str, int, int, str, str, str], Dict[str, Set[str]]] = {}
    for cls in classes:
        if not cls.timeslot:
            continue
        key: Tuple[str, int, int, str, str, str] = (
            cls.timeslot.weekday,
            cls.timeslot.day_index,
            cls.timeslot.sort_order,
            cls.timeslot.label,
            cls.room.code if cls.room else "",
            cls.course.name if cls.course else "",
        )
        entry = grouped.setdefault(
            key,
            {
                "group_codes": set(),
                "program_codes": set(),
            },
        )
        if cls.group and cls.group.code:
            entry["group_codes"].add(cls.group.code)
        if cls.study_program and cls.study_program.code:
            entry["program_codes"].add(cls.study_program.code)

    rows: List[Dict[str, Any]] = []
    for key in sorted(grouped.keys(), key=lambda k: (k[1], k[2], k[3], k[0])):
        weekday, day_index, sort_order, timeslot_label, room_name, course_name = key
        entry = grouped[key]
        rows.append(
            {
                "weekday": weekday,
                "day_index": day_index,
                "sort_order": sort_order,
                "timeslot": timeslot_label,
                "group_code": ", ".join(sorted(entry["group_codes"])),
                "program_code": ", ".join(sorted(entry["program_codes"])),
                "room_name": room_name,
                "course_name": course_name,
            }
        )
    return rows


def _write_teacher_schedule_sheet(
    ws: Worksheet,
    teacher_name: str,
    rows: List[Dict[str, Any]],
    border: Border,
) -> None:
    ws.title = _sanitize_sheet_title(teacher_name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    title_cell = _cell(ws, 1, 1, f"Teacher: {teacher_name}")
    title_cell.font = default_font(bold=True, size=15)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["Day", "Timeslot", "Group", "Program", "Room", "Course"]
    for idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=idx, value=label)
        cell.font = default_font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    ws.column_dimensions["A"].width = float(12)
    ws.column_dimensions["B"].width = float(18)
    ws.column_dimensions["C"].width = float(18)
    ws.column_dimensions["D"].width = float(18)
    ws.column_dimensions["E"].width = float(12)
    ws.column_dimensions["F"].width = float(26)

    if not rows:
        cell = _cell(ws, 3, 1, "No scheduled classes for this teacher.")
        cell.font = default_font(italic=True)
        return

    for idx, row_info in enumerate(rows, start=3):
        ws.cell(row=idx, column=1, value=row_info["weekday"]).border = border
        ws.cell(row=idx, column=2, value=row_info["timeslot"]).border = border
        ws.cell(row=idx, column=3, value=row_info["group_code"]).border = border
        ws.cell(row=idx, column=4, value=row_info["program_code"]).border = border
        ws.cell(row=idx, column=5, value=row_info["room_name"]).border = border
        course_cell = ws.cell(row=idx, column=6, value=row_info["course_name"])
        course_cell.border = border
        course_cell.alignment = Alignment(wrap_text=True, horizontal="left", vertical="center")
        ws.row_dimensions[idx].height = 24


def _build_course_schedule_rows(
    db: Session,
    timetable_id: Optional[int],
    course_id: int,
    group_ids: Optional[List[int]] = None,
    study_program_ids: Optional[List[int]] = None,
) -> List[Dict[str, Any]]:
    query = select(ScheduledClass).options(
        joinedload(ScheduledClass.group),
        joinedload(ScheduledClass.timeslot),
        joinedload(ScheduledClass.room),
        joinedload(ScheduledClass.teacher),
        joinedload(ScheduledClass.course),
        joinedload(ScheduledClass.study_program),
    ).where(
        ScheduledClass.deploy.is_(True),
        ScheduledClass.course_id == course_id,
    )
    if timetable_id is not None:
        query = query.join(Group).where(Group.timetable_id == timetable_id)
    if group_ids is not None:
        query = query.where(ScheduledClass.group_id.in_(group_ids))
    if study_program_ids is not None:
        query = query.where(ScheduledClass.study_program_id.in_(study_program_ids))
    classes = db.execute(query).unique().scalars().all()

    grouped: Dict[Tuple[str, int, int, str, str, str], Dict[str, Set[str]]] = {}
    for cls in classes:
        if not cls.timeslot:
            continue
        key = (
            cls.timeslot.weekday,
            cls.timeslot.day_index,
            cls.timeslot.sort_order,
            cls.timeslot.label,
            cls.teacher.name if cls.teacher else "",
            cls.room.code if cls.room else "",
        )
        entry = grouped.setdefault(
            key,
            {
                "group_codes": set(),
                "program_codes": set(),
            },
        )
        if cls.group and cls.group.code:
            entry["group_codes"].add(cls.group.code)
        if cls.study_program and cls.study_program.code:
            entry["program_codes"].add(cls.study_program.code)

    rows: List[Dict[str, Any]] = []
    for key in sorted(grouped.keys(), key=lambda k: (k[1], k[2], k[3], k[0])):
        weekday, day_index, sort_order, timeslot_label, teacher_name, room_name = key
        entry = grouped[key]
        rows.append(
            {
                "weekday": weekday,
                "day_index": day_index,
                "sort_order": sort_order,
                "timeslot": timeslot_label,
                "teacher_name": teacher_name,
                "group_code": ", ".join(sorted(entry["group_codes"])),
                "program_code": ", ".join(sorted(entry["program_codes"])),
                "room_name": room_name,
            }
        )
    return rows


def _write_course_schedule_sheet(
    ws: Worksheet,
    course_name: str,
    rows: List[Dict[str, Any]],
    border: Border,
) -> None:
    ws.title = _sanitize_sheet_title(course_name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    title_cell = _cell(ws, 1, 1, f"Course: {course_name}")
    title_cell.font = default_font(bold=True, size=15)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["Day", "Timeslot", "Teacher", "Class", "Program", "Room"]
    for idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=idx, value=label)
        cell.font = default_font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    ws.column_dimensions["A"].width = float(12)
    ws.column_dimensions["B"].width = float(18)
    ws.column_dimensions["C"].width = float(20)
    ws.column_dimensions["D"].width = float(20)
    ws.column_dimensions["E"].width = float(18)
    ws.column_dimensions["F"].width = float(12)

    if not rows:
        cell = _cell(ws, 3, 1, "No scheduled classes for this course.")
        cell.font = default_font(italic=True)
        return

    for idx, row_info in enumerate(rows, start=3):
        ws.cell(row=idx, column=1, value=row_info["weekday"]).border = border
        ws.cell(row=idx, column=2, value=row_info["timeslot"]).border = border
        ws.cell(row=idx, column=3, value=row_info["teacher_name"]).border = border
        class_cell = ws.cell(row=idx, column=4, value=row_info["group_code"])
        class_cell.border = border
        room_cell = ws.cell(row=idx, column=5, value=row_info["program_code"])
        room_cell.border = border
        ws.cell(row=idx, column=6, value=row_info["room_name"]).border = border
        class_cell.alignment = Alignment(wrap_text=True, horizontal="left", vertical="center")
        room_cell.alignment = Alignment(wrap_text=True, horizontal="left", vertical="center")
        ws.row_dimensions[idx].height = 24


def _sanitize_sheet_title(title: str, max_length: int = 31) -> str:
    clean = re.sub(r"[\[\]\*:/\\\?']", "", title).strip()
    if not clean:
        return "Sheet"
    return clean[:max_length]


def _logo_path() -> Path:
    return Path(__file__).resolve().parent / "static" / "VGU-Logo.png"


def _insert_logo(ws: Worksheet) -> None:
    logo_file = _logo_path()
    if not logo_file.exists():
        return
    try:
        image = OpenpyxlImage(str(logo_file))
    except ImportError:
        return
    image.width = 120
    image.height = 40
    image.anchor = "A1"
    ws.add_image(image)


def _unique_sheet_titles(names: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    result: List[str] = []
    for name in names:
        title = _sanitize_sheet_title(name)
        if not title:
            title = "Sheet"
        candidate = title
        index = 1
        while candidate in seen:
            index += 1
            candidate = _sanitize_sheet_title(f"{title}-{index}")
            if len(candidate) > 31:
                candidate = candidate[:31]
        seen[candidate] = 1
        result.append(candidate)
    return result


def _write_timetable_sheet(
    db: Session,
    ws: Worksheet,
    payload: Dict[str, Any],
    timetable_id: Optional[int],
    study_program_ids: Optional[List[int]],
    group_ids: Optional[List[int]],
    teacher_ids: Optional[List[int]],
    selected_program_codes: List[str],
    title_fill: PatternFill,
    header_fill: PatternFill,
    lunch_fill: PatternFill,
    blank_fill: PatternFill,
    border: Border,
) -> None:
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False

    groups = payload["groups"]
    timeslots = payload["timeslots"]
    cell_map = payload["cells"]
    program_color_map: Dict[str, str] = {}

    total_cols = 2 + len(groups)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_text = payload.get('cycle_name') or "Foundation Year 2025/26: Semester 2 - Phase 4"
    if selected_program_codes:
        title_text += f" - {', '.join(selected_program_codes)}"
    title_cell = _cell(ws, 1, 1, title_text)
    # title_cell.fill = title_fill
    title_cell.font = default_font(bold=True, size=24)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    subtitle_text = ""
    if selected_program_codes:
        subtitle_text += f" (Study programs: {', '.join(selected_program_codes)})"
    subtitle_cell = _cell(ws, 2, 1, subtitle_text)
    subtitle_cell.alignment = Alignment(horizontal="center")
    subtitle_cell.font = default_font(italic=True)

    has_german_schedule = any(str(ts.get("label", "")).startswith("German") for ts in timeslots)
    if has_german_schedule:
        german_summaries = _german_group_summary(timeslots, cell_map, groups)
        for idx, group in enumerate(groups, start=3):
            summary_cell = _cell(ws, 3, idx, german_summaries.get(group["id"], ""))
            summary_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            summary_cell.font = default_font(bold=False, size=11)
            summary_cell.border = border
        ws.row_dimensions[3].height = 24

    ws["A4"] = "Day"
    ws["B4"] = "Time"
    for cell in [ws["A4"], ws["B4"]]:
        cell.fill = header_fill
        cell.font = default_font(bold=True, size=15)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for idx, group in enumerate(groups, start=3):
        program_codes = ", ".join([p["code"] for p in group.get("programs", [])])
        group_label = group["code"]
        if program_codes:
            group_label += f" ({program_codes})"
        cell = ws.cell(row=4, column=idx, value=group_label)
        cell.fill = header_fill
        cell.font = default_font(bold=True, size=15)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(idx)].width = float(22)
    ws.row_dimensions[4].height = 72

    row = 5
    weekday_names = {1: "MONDAY", 2: "TUESDAY", 3: "WEDNESDAY", 4: "THURSDAY", 5: "FRIDAY"}
    day_to_slots: Dict[str, List[Dict[str, Any]]] = {}
    for ts in timeslots:
        day_to_slots.setdefault(ts["weekday"], []).append(ts)

    active_days: List[Tuple[int, str]] = []
    for day_idx in range(1, 6):
        weekday = weekday_names[day_idx]
        day_slots = day_to_slots.get(weekday, [])
        if not day_slots:
            continue
        has_class = any(
            any(
                cell_map.get(str(slot["id"]), {}).get(str(group["id"]), [])
                for group in groups
            )
            for slot in day_slots
        )
        if has_class:
            active_days.append((day_idx, weekday))

    if not active_days:
        active_days = [(day_idx, weekday_names[day_idx]) for day_idx in range(1, 6) if day_to_slots.get(weekday_names[day_idx])]

    first_active_day = True
    german_rows: List[int] = []
    for day_idx, weekday in active_days:
        day_slots: List[Dict[str, Any]] = sorted(day_to_slots.get(weekday, []), key=lambda item: item["sort_order"])
        if not day_slots:
            continue
        day_start_row = row

        if not first_active_day:
            group_header_row = row
            for idx, group in enumerate(groups, start=3):
                group_label = group["code"]
                cell = ws.cell(row=group_header_row, column=idx, value=group_label)
                cell.fill = header_fill
                cell.font = default_font(bold=True, size=15)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            ws.cell(row=group_header_row, column=2).border = border
            ws.row_dimensions[group_header_row].height = 72
            row += 1

        first_active_day = False

        prev_german_row: Optional[int] = None
        prev_german_slot_items: Optional[Dict[int, Tuple[str, str]]] = None
        for slot_index, slot in enumerate(day_slots):
            slot_items = cell_map.get(str(slot["id"]), {})
            next_slot = day_slots[slot_index + 1] if slot_index + 1 < len(day_slots) else None
            is_german_slot = str(slot.get("label", "")).startswith("German")
            if is_german_slot:
                slot_start_row = row
                required_row = row
                slot_end_row = row
                ws.merge_cells(start_row=slot_start_row, start_column=2, end_row=slot_end_row, end_column=2)
                time_cell = ws.cell(row=slot_start_row, column=2, value=slot["label"])
                time_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                time_cell.border = border

                row_height = _OVERLAY_MIN_HEIGHT
                current_items: Dict[int, Tuple[str, str]] = {}
                visible_items: Dict[int, Tuple[str, str]] = {}
                for col_idx2, group in enumerate(groups, start=3):
                    items = slot_items.get(str(group["id"]), [])
                    cell = _cell(ws, required_row, col_idx2)
                    if not items:
                        cell.fill = blank_fill
                        cell.border = border
                        continue
                    text = "\n".join(_format_german_item(item) for item in items)
                    first_item = items[0]
                    fill_color = _get_fill_color(first_item)
                    current_items[group["id"]] = (text, fill_color)
                    if prev_german_slot_items is not None and prev_german_slot_items.get(group["id"]) == current_items[group["id"]]:
                        cell.fill = blank_fill
                        cell.border = border
                        continue

                    visible_items[group["id"]] = (text, fill_color)
                    cell.value = text
                    cell.fill = PatternFill("solid", fgColor=fill_color)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.font = default_font(bold=True, size=15)
                    cell.border = border
                    column_width = _column_range_width_chars(ws, col_idx2, col_idx2)
                    needed_height = _row_height_for_text(text, width_cols=column_width, min_height=_OVERLAY_MIN_HEIGHT, font_size=15)
                    row_height = max(row_height, needed_height)

                segment_start: Optional[int] = None
                segment_text: Optional[str] = None
                segment_fill: Optional[str] = None
                for col_idx2, group in enumerate(groups, start=3):
                    entry = visible_items.get(group["id"])
                    if entry is None:
                        if segment_start is not None:
                            segment_end = col_idx2 - 1
                            if segment_end > segment_start:
                                assert segment_text is not None and segment_fill is not None
                                merged_width = _column_range_width_chars(ws, segment_start, segment_end)
                                row_height = max(
                                    row_height,
                                    _row_height_for_text(
                                        segment_text,
                                        width_cols=merged_width,
                                        min_height=_OVERLAY_MIN_HEIGHT,
                                        font_size=15,
                                    ),
                                )
                                ws.merge_cells(
                                    start_row=required_row,
                                    start_column=segment_start,
                                    end_row=required_row,
                                    end_column=segment_end,
                                )
                                merged_cell = _cell(ws, required_row, segment_start)
                                merged_cell.value = segment_text
                                merged_cell.fill = PatternFill("solid", fgColor=segment_fill)
                                merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                merged_cell.font = default_font(bold=True, size=15)
                                for c in range(segment_start, segment_end + 1):
                                    ws.cell(row=required_row, column=c).border = border
                            segment_start = None
                            segment_text = None
                            segment_fill = None
                        continue

                    text, fill_color = entry
                    if segment_start is None:
                        segment_start = col_idx2
                        segment_text = text
                        segment_fill = fill_color
                        continue
                    if text == segment_text and fill_color == segment_fill:
                        continue

                    segment_end = col_idx2 - 1
                    if segment_end > segment_start:
                        assert segment_text is not None and segment_fill is not None
                        merged_width = _column_range_width_chars(ws, segment_start, segment_end)
                        row_height = max(
                            row_height,
                            _row_height_for_text(
                                segment_text,
                                width_cols=merged_width,
                                min_height=_OVERLAY_MIN_HEIGHT,
                                font_size=15,
                            ),
                        )
                        ws.merge_cells(
                            start_row=required_row,
                            start_column=segment_start,
                            end_row=required_row,
                            end_column=segment_end,
                        )
                        merged_cell = _cell(ws, required_row, segment_start)
                        merged_cell.value = segment_text
                        merged_cell.fill = PatternFill("solid", fgColor=segment_fill)
                        merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        merged_cell.font = default_font(bold=True, size=15)
                        for c in range(segment_start, segment_end + 1):
                            ws.cell(row=required_row, column=c).border = border
                    segment_start = col_idx2
                    segment_text = text
                    segment_fill = fill_color

                if segment_start is not None:
                    segment_end = len(groups) + 2
                    if segment_end > segment_start:
                        assert segment_text is not None and segment_fill is not None
                        merged_width = _column_range_width_chars(ws, segment_start, segment_end)
                        row_height = max(
                            row_height,
                            _row_height_for_text(
                                segment_text,
                                width_cols=merged_width,
                                min_height=_OVERLAY_MIN_HEIGHT,
                                font_size=15,
                            ),
                        )
                        ws.merge_cells(
                            start_row=required_row,
                            start_column=segment_start,
                            end_row=required_row,
                            end_column=segment_end,
                        )
                        merged_cell = _cell(ws, required_row, segment_start)
                        merged_cell.value = segment_text
                        merged_cell.fill = PatternFill("solid", fgColor=segment_fill)
                        merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        merged_cell.font = default_font(bold=True, size=15)
                        for c in range(segment_start, segment_end + 1):
                            ws.cell(row=required_row, column=c).border = border

                if prev_german_slot_items is not None and prev_german_row is not None:
                    for col_idx2, group in enumerate(groups, start=3):
                        prev_item = prev_german_slot_items.get(group["id"])
                        curr_item = current_items.get(group["id"])
                        if prev_item is not None and curr_item is not None and prev_item == curr_item:
                            ws.merge_cells(
                                start_row=prev_german_row,
                                start_column=col_idx2,
                                end_row=required_row,
                                end_column=col_idx2,
                            )
                            merged_cell = _cell(ws, prev_german_row, col_idx2)
                            merged_cell.value = curr_item[0]
                            merged_cell.fill = PatternFill("solid", fgColor=curr_item[1])
                            merged_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            merged_cell.font = default_font(bold=True, size=15)
                            for r in range(prev_german_row, required_row + 1):
                                cell = _cell(ws, r, col_idx2)
                                cell.border = border
                                if r != prev_german_row:
                                    cell.fill = blank_fill

                ws.row_dimensions[required_row].height = row_height
                german_rows.append(required_row)
                row += 1

                current_end = _parse_time_minutes(slot.get("end"))
                next_start = _parse_time_minutes(next_slot.get("start") if next_slot else None)
                lunch_needed = False
                if current_end is not None and next_start is not None and current_end <= 12 * 60 and next_start >= 13 * 60:
                    lunch_row = row
                    ws.merge_cells(start_row=lunch_row, start_column=2, end_row=lunch_row, end_column=total_cols)
                    lunch = ws.cell(row=lunch_row, column=2, value="Lunch break")
                    lunch.fill = lunch_fill
                    lunch.alignment = Alignment(horizontal="center")
                    lunch.font = default_font(italic=True)
                    for c in range(2, total_cols + 1):
                        ws.cell(row=lunch_row, column=c).border = border
                    row += 1
                    prev_german_slot_items = None
                    prev_german_row = None
                else:
                    prev_german_slot_items = current_items
                    prev_german_row = required_row
                continue

            overlay_rows_content: List[List[Dict[str, Any]]] = []
            for group in groups:
                items = _overlay_items(slot_items.get(str(group["id"]), []))
                deduped: List[Dict[str, Any]] = []
                for item in items:
                    if not any(_same_overlay(item, d) for d in deduped):
                        deduped.append(item)
                overlay_rows_content.append(deduped)

            unique_overlays: List[Dict[str, Any]] = []
            for overlays in overlay_rows_content:
                for item in overlays:
                    found = next((uo for uo in unique_overlays if _same_overlay_connectable(item, uo)), None)
                    if found:
                        found["program_codes"] = sorted(set(found.get("program_codes", []) + item.get("program_codes", [])))
                        found["group_codes"] = sorted(set(found.get("group_codes", []) + item.get("group_codes", [])))
                        found["all_group"] = found.get("all_group", False) or item.get("all_group", False)
                    elif not any(_same_overlay(item, uo) for uo in unique_overlays):
                        unique_overlays.append(
                            {
                                **item,
                                "program_codes": list(item.get("program_codes", [])),
                                "group_codes": list(item.get("group_codes", [])),
                                "all_group": item.get("all_group", False),
                            }
                        )

            kind_order = {"required": 0, "program": 1, "elective": 2}
            unique_overlays.sort(
                key=lambda item: (
                    kind_order.get(str(item.get("kind") or ""), 3),
                    item.get("course_name", ""),
                    item.get("teacher_name", ""),
                    item.get("room_name", ""),
                )
            )
            _assign_program_colors_for_slot(unique_overlays, program_color_map)
            overlay_depth = len(unique_overlays)
            slot_start_row = row
            required_row = row
            overlay_rows = [required_row + i + 1 for i in range(overlay_depth)]
            slot_end_row = required_row + overlay_depth

            ws.merge_cells(start_row=slot_start_row, start_column=2, end_row=slot_end_row, end_column=2)
            time_cell = ws.cell(row=slot_start_row, column=2, value=slot["label"])
            time_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            time_cell.border = border

            required_row_height = REQUIRED_ROW_HEIGHT
            ws.row_dimensions[required_row].height = required_row_height

            col_idx = 3
            while col_idx <= total_cols:
                group_index = col_idx - 3
                if group_index >= len(groups):
                    break
                group = groups[group_index]
                items = slot_items.get(str(group["id"]), [])
                req = _required_item(items)
                if not req:
                    col_idx += 1
                    continue
                end_col = col_idx
                next_col = col_idx + 1
                while next_col <= total_cols:
                    next_group = groups[next_col - 3]
                    next_items = slot_items.get(str(next_group["id"]), [])
                    next_req = _required_item(next_items)
                    if not next_req or not _same_overlay(req, next_req):
                        break
                    end_col = next_col
                    next_col += 1
                if end_col > col_idx:
                    _write_overlay_merge(ws, required_row, col_idx, end_col, req, border)
                    req_width = _column_range_width_chars(ws, col_idx, end_col)
                    req_height = _row_height_for_text(
                        _format_item(req),
                        req_width,
                        min_height=REQUIRED_ROW_HEIGHT,
                        font_size=15,
                    )
                    required_row_height = max(required_row_height, req_height)
                    col_idx = end_col + 1
                else:
                    req_cell = _cell(ws, required_row, col_idx)
                    req_cell.value = _format_item(req)
                    req_cell.fill = PatternFill("solid", fgColor=FILL_MAP.get(req["color_key"], "D9D2E9"))
                    req_cell.border = border
                    req_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    req_cell.font = default_font(bold=True, size=15)
                    req_width = _column_range_width_chars(ws, col_idx, col_idx)
                    req_height = _row_height_for_text(
                        _format_item(req),
                        req_width,
                        min_height=REQUIRED_ROW_HEIGHT,
                        font_size=15,
                    )
                    required_row_height = max(required_row_height, req_height)
                    col_idx += 1
            ws.row_dimensions[required_row].height = required_row_height

            for col_idx2, group in enumerate(groups, start=3):
                items = slot_items.get(str(group["id"]), [])
                req = _required_item(items)
                req_cell = _cell(ws, required_row, col_idx2)
                if not req_cell.value:
                    req_cell.fill = blank_fill
                    req_cell.border = border

                for overlay_row in overlay_rows:
                    cell = _cell(ws, overlay_row, col_idx2)
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.font = default_font(bold=True, size=15)

            merged_overlay_columns: Dict[int, Set[int]] = {}
            for overlay_idx, overlay_item in enumerate(unique_overlays):
                overlay_row = required_row + 1 + overlay_idx
                row_has_overlay = [any(_same_overlay_connectable(item, overlay_item) for item in group_items) for group_items in overlay_rows_content]
                present_cols = [3 + idx for idx, has in enumerate(row_has_overlay) if has]
                # Compute height based on each contiguous merged segment, not on total present columns.
                needed_height = _OVERLAY_MIN_HEIGHT
                segment_width = 0
                segment_start: Optional[int] = None
                for idx, has in enumerate(row_has_overlay + [False]):
                    if has:
                        if segment_start is None:
                            segment_start = idx
                        segment_width += 1
                        continue
                    if segment_width and segment_start is not None:
                        start_col = 3 + segment_start
                        end_col = 3 + segment_start + segment_width - 1
                        segment_chars = _column_range_width_chars(ws, start_col, end_col)
                        segment_height = _row_height_for_text(
                            _format_item(overlay_item),
                            segment_chars,
                            min_height=_OVERLAY_MIN_HEIGHT,
                            font_size=15,
                        )
                        needed_height = max(needed_height, segment_height)
                        segment_width = 0
                        segment_start = None
                row_dimension: RowDimension = ws.row_dimensions[overlay_row]
                current_height = float(getattr(row_dimension, "height", 0) or 0)
                setattr(row_dimension, "height", float(max(current_height, needed_height)))
                merged_overlay_columns[overlay_row] = set()
                if overlay_item.get("kind") == "program" and len(present_cols) > 1:
                    start_col = min(present_cols)
                    end_col = max(present_cols)
                    merged_overlay_columns[overlay_row].update(range(start_col, end_col + 1))
                    _write_overlay_merge(ws, overlay_row, start_col, end_col, overlay_item, border)
                else:
                    col_idx = 3
                    while col_idx <= total_cols:
                        if col_idx - 3 >= len(row_has_overlay):
                            break
                        if not row_has_overlay[col_idx - 3]:
                            col_idx += 1
                            continue
                        end_col = col_idx
                        next_col = col_idx + 1
                        while next_col <= total_cols and next_col - 3 < len(row_has_overlay) and row_has_overlay[next_col - 3]:
                            end_col = next_col
                            next_col += 1
                        merged_overlay_columns[overlay_row].update(range(col_idx, end_col + 1))
                        _write_overlay_merge(ws, overlay_row, col_idx, end_col, overlay_item, border)
                        col_idx = end_col + 1

            for overlay_idx, overlay_item in enumerate(unique_overlays):
                overlay_row = required_row + 1 + overlay_idx
                for col_idx2, group in enumerate(groups, start=3):
                    if col_idx2 in merged_overlay_columns.get(overlay_row, set()):
                        continue
                    group_items = overlay_rows_content[col_idx2 - 3]
                    if not any(_same_overlay(item, overlay_item) for item in group_items):
                        cell = _cell(ws, overlay_row, col_idx2)
                        cell.fill = blank_fill
                        cell.border = border

            for rr in range(slot_start_row, slot_end_row + 1):
                ws.cell(row=rr, column=2).border = border

            row = slot_end_row + 1

            next_start = _parse_time_minutes(next_slot.get("start") if next_slot else None)
            slot_end = _parse_time_minutes(slot.get("end"))
            lunch_needed = False
            if slot_end is not None and next_start is not None:
                lunch_needed = slot_end <= 12 * 60 and next_start >= 13 * 60
            if lunch_needed:
                lunch_row = row
                ws.merge_cells(start_row=lunch_row, start_column=2, end_row=lunch_row, end_column=total_cols)
                lunch = ws.cell(row=lunch_row, column=2, value="Lunch break")
                lunch.fill = lunch_fill
                lunch.alignment = Alignment(horizontal="center")
                lunch.font = default_font(italic=True)
                for c in range(2, total_cols + 1):
                    ws.cell(row=lunch_row, column=c).border = border
                row += 1

        day_end_row = row - 1
        day_separator_border = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thick", color="000000"),
        )
        for c in range(1, total_cols + 1):
            cell = _cell(ws, day_end_row, c)
            cell.border = day_separator_border
        ws.merge_cells(start_row=day_start_row, start_column=1, end_row=day_end_row, end_column=1)
        dcell = ws.cell(row=day_start_row, column=1, value=weekday)
        dcell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        dcell.font = default_font(bold=True)
        dcell.fill = PatternFill("solid", fgColor="E2EFDA")
        dcell.border = border
        for rr in range(day_start_row, day_end_row + 1):
            ws.cell(row=rr, column=1).border = border

    timetable_end_row = row - 1
    last_group_col = 2 + len(groups)
    max_col = ws.max_column

    if german_rows:
        max_german_height = max(
            float(getattr(ws.row_dimensions[r], "height", 0) or 0)
            for r in german_rows
        )
        if max_german_height > 0:
            for r in german_rows:
                ws.row_dimensions[r].height = max_german_height

    white_fill = PatternFill("solid", fgColor="FFFFFF")
    for r in range(5, timetable_end_row + 1):
        for c in range(last_group_col + 1, max_col + 1):
            cell = _cell(ws, r, c)
            cell.value = None
            cell.fill = white_fill
            cell.border = Border()

    start_teacher = row + 2
    ws.cell(row=start_teacher, column=1, value="Teacher load summary").font = default_font(bold=True, size=12)
    ws.cell(row=start_teacher + 1, column=1, value="Teacher").font = default_font(bold=True)
    ws.cell(row=start_teacher + 1, column=2, value="Taught timeslots").font = default_font(bold=True)
    for idx, item in enumerate(teacher_load_rows(db, timetable_id, study_program_ids, group_ids=group_ids, teacher_ids=teacher_ids), start=start_teacher + 2):
        ws.cell(row=idx, column=1, value=item["teacher"])
        ws.cell(row=idx, column=2, value=item["timeslot_count"])

    ws.column_dimensions["A"].width = float(10)
    ws.column_dimensions["B"].width = float(16)


def export_timetable_xlsx(
    db: Session,
    output_path: str | Path,
    timetable_id: Optional[int] = None,
    study_program_ids: Optional[List[int]] = None,
    group_ids: Optional[List[int]] = None,
    group_tag_ids: Optional[List[int]] = None,
    teacher_ids: Optional[List[int]] = None,
    course_ids: Optional[List[int]] = None,
) -> Path:
    output_path = Path(output_path)
    title_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    lunch_fill = PatternFill("solid", fgColor="F2F2F2")
    blank_fill = PatternFill("solid", fgColor=BLANK_FILL)
    thin_gray = Side(style="thin", color="444444")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    selected_program_codes: List[str] = []
    if study_program_ids is not None:
        selected_program_codes = [str(prog.code) for prog in db.scalars(select(StudyProgram).where(StudyProgram.id.in_(study_program_ids))).all()]

    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    payload = build_timetable_payload(
        db,
        timetable_id,
        study_program_ids=study_program_ids,
        group_ids=group_ids,
        teacher_ids=teacher_ids,
    )

    if course_ids is not None:
        from .models import Course
        course_query = select(Course).where(Course.id.in_(course_ids)).order_by(Course.name)
        if timetable_id is not None:
            course_query = (
                course_query.join(ScheduledClass, ScheduledClass.course_id == Course.id)
                .join(Group, Group.id == ScheduledClass.group_id)
                .where(ScheduledClass.deploy.is_(True), Group.timetable_id == timetable_id)
                .distinct()
            )
        course_sheets = list(db.scalars(course_query).all())
        sheet_titles = _unique_sheet_titles([str(c.name or c.id) for c in course_sheets])
        for idx, (course, sheet_title) in enumerate(zip(course_sheets, sheet_titles)):
            if idx == 0:
                ws.title = sheet_title
            else:
                ws = wb.create_sheet(title=sheet_title)
            rows = _build_course_schedule_rows(db, timetable_id, cast(int, getattr(course, 'id')), group_ids=group_ids, study_program_ids=study_program_ids)
            _write_course_schedule_sheet(ws, str(course.name), rows, border)
            _insert_logo(ws)
        wb.save(output_path)
        return output_path

    if teacher_ids is not None:
        from .models import Teacher
        teacher_query = select(Teacher).where(Teacher.id.in_(teacher_ids)).order_by(Teacher.name)
        if timetable_id is not None:
            teacher_query = (
                teacher_query.join(ScheduledClass, ScheduledClass.teacher_id == Teacher.id)
                .join(Group, Group.id == ScheduledClass.group_id)
                .where(ScheduledClass.deploy.is_(True), Group.timetable_id == timetable_id)
                .distinct()
            )
        teacher_sheets = list(db.scalars(teacher_query).all())
        sheet_titles = _unique_sheet_titles([str(t.name or t.id) for t in teacher_sheets])
        for idx, (teacher, sheet_title) in enumerate(zip(teacher_sheets, sheet_titles)):
            if idx == 0:
                ws.title = sheet_title
            else:
                ws = wb.create_sheet(title=sheet_title)
            rows = _build_teacher_schedule_rows(db, timetable_id, cast(int, getattr(teacher, 'id')))
            _write_teacher_schedule_sheet(ws, str(teacher.name), rows, border)
            _insert_logo(ws)
        wb.save(output_path)
        return output_path

    if group_ids is not None and group_tag_ids is None:
        group_sheets = payload["groups"]
        program_color_map = _build_program_color_map(payload)
        sheet_titles = _unique_sheet_titles([str(group["code"]) for group in group_sheets])
        for idx, (group, sheet_title) in enumerate(zip(group_sheets, sheet_titles)):
            if idx == 0:
                ws.title = sheet_title
            else:
                ws = wb.create_sheet(title=sheet_title)
            _write_group_timetable_sheet(ws, payload, group, blank_fill, border, program_color_map)
            _insert_logo(ws)
        wb.save(output_path)
        return output_path

    ws.title = "Phase4 Timetable"
    _write_timetable_sheet(
        db,
        ws,
        payload,
        timetable_id,
        study_program_ids,
        group_ids,
        teacher_ids,
        selected_program_codes,
        title_fill,
        header_fill,
        lunch_fill,
        blank_fill,
        border,
    )
    _insert_logo(ws)

    filtered_export = study_program_ids is not None or group_ids is not None or teacher_ids is not None
    program_sheets: List[StudyProgram] = []
    if not filtered_export:
        program_sheets = list(
            db.scalars(
                select(StudyProgram)
                .join(GroupStudyProgram)
                .join(Group)
                .where(Group.timetable_id == timetable_id)
                .distinct()
            ).all()
        ) if timetable_id is not None else list(db.scalars(select(StudyProgram)).all())

    if program_sheets:
        sheet_titles = _unique_sheet_titles([str(prog.code) for prog in program_sheets])
        for prog, sheet_title in zip(program_sheets, sheet_titles):
            if sheet_title == ws.title:
                sheet_title = _sanitize_sheet_title(f"{sheet_title}-1")
            program_ws = wb.create_sheet(title=sheet_title)
            prog_id = cast(int, prog.id)
            prog_payload = build_timetable_payload(db, timetable_id, study_program_ids=[prog_id])
            _write_timetable_sheet(
                db,
                program_ws,
                prog_payload,
                timetable_id,
                [prog_id],
                None,
                None,
                [str(prog.code)],
                title_fill,
                header_fill,
                lunch_fill,
                blank_fill,
                border,
            )
            _insert_logo(program_ws)

    wb.save(output_path)
    return output_path