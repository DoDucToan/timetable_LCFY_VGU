"""Excel export matching the selected program filter in the Group schedule dialog.

Shares the /api/bootstrap payload produced by app.main so that display and
export use the same deployed classes, timetable, group and note values.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def build_group_schedule_rows(
    payload: Dict[str, Any],
    group_id: int,
    program_ids: Optional[Sequence[int]] = None,
    include_electives: bool = True,
) -> List[Dict[str, str]]:
    """Mirror JS buildGroupScheduleRows: filtering, merging, and sort order.

    Program filter only affects program-specific classes. A required or
    elective class belongs to the whole group. The elective toggle removes
    elective classes even if they are included at the group level.
    """
    group = next((g for g in payload.get("groups", []) if str(g.get("id")) == str(group_id)), None)
    if group is None:
        raise ValueError("Group does not belong to this timetable")

    selected_ids = {int(p) for p in (program_ids or [])}
    selected_codes = {
        str(p.get("code") or "").strip().upper()
        for p in group.get("programs", [])
        if int(p.get("id") or 0) in selected_ids
    }
    selected_codes.discard("")
    timeslots = {str(slot.get("id")): slot for slot in payload.get("timeslots", [])}
    grouped: Dict[tuple, Dict[str, Any]] = {}
    for slot_id, groups in (payload.get("cells") or {}).items():
        slot = timeslots.get(str(slot_id))
        if not slot or not isinstance(groups, dict):
            continue
        for item in groups.get(str(group_id), []):
            kind = str(item.get("kind") or "")
            if not include_electives and kind == "elective":
                continue
            codes = [str(code) for code in (item.get("program_codes") or []) if code]
            if selected_codes and kind not in {"required", "elective"}:
                if not any(code.strip().upper() in selected_codes for code in codes):
                    continue
            day = str(slot.get("weekday") or "")
            order = int(slot.get("sort_order") or 0)
            time_label = str(slot.get("label") or "")
            teacher = str(item.get("teacher_name") or "")
            room = str(item.get("room_name") or "")
            course = str(item.get("course_name") or "")
            key = (day, order, time_label, teacher, room, kind, course)
            if key not in grouped:
                grouped[key] = {
                    "weekday": day,
                    "sort_order": order,
                    "timeslot": time_label,
                    "teacher_name": teacher,
                    "room_name": room,
                    "course_name": f"(Elective) {course}" if kind == "elective" else course,
                    "kind": kind,
                    "program_codes": set(),
                    "notes_set": set(),
                }
            grouped[key]["program_codes"].update(codes)
            note = str(item.get("notes") or "").strip()
            if note:
                grouped[key]["notes_set"].add(note)

    result = []
    for row in grouped.values():
        result.append({
            "weekday": row["weekday"],
            "sort_order": row["sort_order"],
            "timeslot": row["timeslot"],
            "teacher_name": row["teacher_name"],
            "program_code": ", ".join(sorted(row["program_codes"])),
            "room_name": row["room_name"],
            "course_name": row["course_name"],
            "notes": " | ".join(sorted(row["notes_set"])),
            "kind": row["kind"],
        })
    return sorted(result, key=lambda r: (r["sort_order"], r["weekday"], r["timeslot"]))


def write_group_schedule_xlsx(
    payload: Dict[str, Any],
    output_path: str | Path,
    group_id: int,
    program_ids: Optional[Sequence[int]] = None,
    include_electives: bool = True,
) -> int:
    """Write a single-sheet Excel table with exactly the filtered modal rows.

    Returns the number of exported classes. Prevents Excel formula injection
    by marking all user-supplied values as literal strings.
    """
    group = next((g for g in payload.get("groups", []) if str(g.get("id")) == str(group_id)), None)
    if group is None:
        raise ValueError("Group does not belong to this timetable")
    rows = build_group_schedule_rows(payload, group_id, program_ids, include_electives)
    selected = {int(x) for x in (program_ids or [])}
    codes = sorted(
        str(p.get("code") or "") for p in group.get("programs", [])
        if int(p.get("id") or 0) in selected
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Group schedule"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C6"
    widths = (17, 23, 28, 24, 18, 43, 50)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    navy = "1F4E78"
    pale = "EAF1FB"
    yellow = "FFF6BD"
    line = Side(style="thin", color="CED9E7")
    border = Border(bottom=line, left=line, right=line, top=line)

    ws.merge_cells("A1:G1")
    title = ws["A1"]
    title.value = f"Group schedule: {group.get('code') or group_id}"
    title.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    for row_num, caption in (
        (2, "Programs: " + (", ".join(codes) if codes else "All programs")),
        (3, "Electives: " + ("Included" if include_electives else "Hidden")),
    ):
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=7)
        cell = ws.cell(row=row_num, column=1, value=caption)
        cell.data_type = "s"
        cell.font = Font(name="Calibri", size=11, bold=row_num == 2, color="31445A")
        cell.alignment = Alignment(vertical="center", indent=1)
        cell.fill = PatternFill("solid", fgColor=pale)
        ws.row_dimensions[row_num].height = 23

    headers = ("Day", "Timeslot", "Teacher", "Program", "Room", "Course", "Notes")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[5].height = 26
    fields = ("weekday", "timeslot", "teacher_name", "program_code", "room_name", "course_name", "notes")
    for row_no, data in enumerate(rows, 6):
        is_elective = data.get("kind") == "elective"
        shade = yellow if is_elective else ("F4F8FD" if row_no % 2 == 0 else "FFFFFF")
        for col, key in enumerate(fields, 1):
            cell = ws.cell(row=row_no, column=col, value=str(data.get(key) or ""))
            cell.data_type = "s"  # protect against formula-like content
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=shade)
            cell.font = Font(name="Calibri", size=11, color="25364A")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        longest = max(len(data.get("course_name", "")), len(data.get("notes", "")))
        ws.row_dimensions[row_no].height = min(100, max(28, 16 * (1 + longest // 48)))

    ws.auto_filter.ref = f"A5:G{max(5, 5 + len(rows))}"
    wb.save(Path(output_path))
    return len(rows)


# LCFY_GROUP_PROGRAM_EXPORT_PATCH_V4
